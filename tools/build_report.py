"""Build report/HeartSense_MLOps_Assignment1_Report.docx from the repository artefacts.

python tools/build_report.py --author "Name" --roll "2025XXXXXXX" --repo URL --video URL
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots"
FIG = ROOT / "reports" / "figures"
OUT = ROOT / "report" / "HeartSense_MLOps_Assignment1_Report.docx"


def style_doc(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for s, size, color in (
        ("Heading 1", 16, "1F5F8B"),
        ("Heading 2", 13, "1F5F8B"),
        ("Heading 3", 11.5, "333333"),
    ):
        doc.styles[s].font.size = Pt(size)
        doc.styles[s].font.color.rgb = RGBColor.from_string(color)
        doc.styles[s].font.name = "Calibri"
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2.0)


def para(doc, text, bold=False, italic=False, size=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            r = p.add_run(it[0])
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(2)


def code(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(6)


def figure(doc, path: Path, caption: str, width_cm: float = 15.5):
    if not path.exists():
        para(doc, f"[missing figure: {path.relative_to(ROOT)}]", italic=True)
        return
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, caption, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)


def table(doc, header, rows, col_widths=None, font_size=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold = True
        r.font.size = Pt(font_size)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(font_size)
    if col_widths:
        for row in t.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--author", default="Mahamaya Mishra")
    ap.add_argument("--roll", default="<roll number>")
    ap.add_argument(
        "--repo", default="https://github.com/mahamaya-mishra/heartsense-mlops-assignment"
    )
    ap.add_argument("--video", default="<Google Drive link to the video>")
    a = ap.parse_args()

    meta = json.loads((ROOT / "model_store" / "metadata.json").read_text())
    eda = json.loads((ROOT / "reports" / "eda_profile.json").read_text())
    board = pd.read_csv(ROOT / "reports" / "leaderboard.csv")
    m = meta["metrics"]

    doc = Document()
    style_doc(doc)

    # ---------------------------------------------------------------- title page
    for _ in range(5):
        doc.add_paragraph()
    para(doc, "HeartSense", bold=True, size=30, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(
        doc,
        "Heart-disease risk prediction with an end-to-end MLOps pipeline",
        size=15,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=24,
    )
    para(
        doc,
        "Machine Learning Operations (AIMLCZG523)  -  Assignment 1",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    para(
        doc,
        "BITS Pilani, Work Integrated Learning Programmes",
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=30,
    )
    para(doc, f"{a.author}", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, f"Roll number: {a.roll}", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
    para(
        doc, f"Code repository: {a.repo}", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2
    )
    para(
        doc,
        f"Video walkthrough: {a.video}",
        size=10.5,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=2,
    )
    para(
        doc,
        f"Model: {meta['model_version']}  -  trained {meta['trained_at'][:10]}  -  git {meta['git_sha']}",
        size=9.5,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    doc.add_page_break()

    # ---------------------------------------------------------------- 1. overview
    doc.add_heading("1. Problem statement and solution overview", 1)
    para(
        doc,
        "Coronary heart disease is the leading cause of death worldwide, and the first clinical "
        "contact usually produces a small set of routine measurements: age, blood pressure, cholesterol, "
        "resting and exercise ECG findings, exercise tolerance and a thallium perfusion result. The task "
        "is to turn those 13 measurements into a calibrated probability of heart disease and to deliver "
        "that model the way a production team would: reproducible, tested, containerised, deployed on "
        "Kubernetes, monitored, and rebuilt automatically on every change.",
    )
    para(
        doc,
        "HeartSense is the resulting system. Everything is code in one repository and runs from a clean "
        "checkout with a pinned requirements file. The pipeline has a build-time lane (data, EDA, "
        "training with MLflow tracking, CI/CD that lints, tests, trains and publishes a container image) "
        "and a run-time lane (Helm-deployed FastAPI service on Kubernetes behind a LoadBalancer, scraped "
        "by Prometheus and visualised in Grafana).",
    )
    figure(
        doc,
        ROOT / "docs" / "architecture.png",
        "Figure 1. End-to-end architecture: build time (top) and run time (bottom).",
        16.5,
    )
    doc.add_heading("Key results", 2)
    bullets(
        doc,
        [
            (
                "Champion model: ",
                f"{meta['family'].replace('_', ' ')} selected on mean cross-validated ROC-AUC "
                f"{m['cv_roc_auc_mean']:.3f} ± {m['cv_roc_auc_std']:.3f} (5-fold x 2 repeats).",
            ),
            (
                "Hold-out test (61 patients): ",
                f"ROC-AUC {m['test_roc_auc']:.3f}, accuracy {m['test_accuracy']:.3f}, "
                f"recall {m['test_recall']:.3f}, precision {m['test_precision']:.3f}, "
                f"{m['test_false_negatives']} false negatives.",
            ),
            (
                "Quality gates: ",
                "ruff lint and format, 56 pytest tests (fast suite in about 2 s, plus an end-to-end "
                "training test), Helm lint and kubeconform, Docker smoke test; every push runs them.",
            ),
            (
                "Operations: ",
                "two replicas with probes, HPA and PDB; one JSON log line per request without patient "
                "data; 14-panel Grafana dashboard and five Prometheus alert rules including a prediction-drift check.",
            ),
        ],
    )

    # ---------------------------------------------------------------- 2. setup
    doc.add_heading("2. Setup and installation", 1)
    para(
        doc,
        "Requirements: Python 3.11 or 3.12, Git, Docker (Docker Desktop or colima). minikube, kubectl and "
        "Helm are only needed for the deployment step. All commands are also wrapped as Makefile targets "
        "(`make help`).",
    )
    code(
        doc,
        "git clone " + a.repo + ".git && cd heartsense-mlops-assignment\n"
        "python3.12 -m venv .venv && source .venv/bin/activate\n"
        "pip install -r requirements.txt            # pinned; installs the heartsense package too\n"
        "export MLFLOW_DISABLE_AGENT_HINT=1 PYTHONPATH=src\n\n"
        "python -m heartsense.ingest                # download + SHA-256 verify + clean\n"
        "python -m heartsense.explore               # EDA figures and profile\n"
        "python -m heartsense.train                 # tune 5 families, MLflow, registry, export\n"
        "make lint && pytest --cov                  # quality gates\n"
        "docker build -t heartsense-api:1.0.0 . && docker run -d -p 8000:8000 heartsense-api:1.0.0\n"
        "bash scripts/smoke_test.sh http://127.0.0.1:8000\n\n"
        "bash scripts/minikube_up.sh                # local Kubernetes\n"
        "IMAGE=heartsense-api:1.0.0 bash scripts/deploy.sh   # Helm release + Prometheus + Grafana\n"
        "bash scripts/port_forward.sh && bash scripts/smoke_test.sh http://127.0.0.1:8000",
    )
    table(
        doc,
        ["Layer", "Choice", "Why"],
        [
            [
                "Language / ML",
                "Python 3.12, scikit-learn 1.9, pandas 3",
                "Mature tabular stack; one Pipeline object for train and serve",
            ],
            [
                "Tracking",
                "MLflow 3.17 (SQLite + local artifacts)",
                "Nested runs, model signature, registry with aliases",
            ],
            [
                "API",
                "FastAPI + pydantic v2 + uvicorn",
                "Typed validation, OpenAPI docs, async-friendly",
            ],
            ["Tests / lint", "pytest, pytest-cov, ruff", "Fast, standard, CI friendly"],
            [
                "CI/CD",
                "GitHub Actions",
                "Lint -> tests (2 Python versions) -> train -> image build + smoke test -> GHCR",
            ],
            [
                "Container",
                "python:3.12-slim multi-stage, non-root",
                "Small attack surface, healthcheck",
            ],
            [
                "Deployment",
                "Helm chart on minikube (Kubernetes 1.37)",
                "Deployment, LoadBalancer Service, HPA, PDB, probes",
            ],
            [
                "Monitoring",
                "Prometheus 3.5 + Grafana 12",
                "Pod discovery via annotations, provisioned dashboard, alert rules",
            ],
        ],
        [3.2, 5.6, 7.7],
    )

    # ---------------------------------------------------------------- 3. data + EDA
    doc.add_heading("3. Data acquisition and exploratory data analysis", 1)
    para(
        doc,
        "The UCI Heart Disease dataset (Cleveland Clinic subset, id 45, CC BY 4.0) has 303 patients and 13 "
        "clinical features. `heartsense.ingest` downloads `processed.cleveland.data`, verifies its SHA-256 "
        f"hash ({meta['dataset']['sha256'][:16]}...), renames the UCI abbreviations to descriptive names, "
        "binarises the 0-4 diagnosis into a target (1 = disease), coerces invalid category codes and "
        "physiologically implausible readings to missing, and drops duplicates. Missing values are kept as "
        "NaN on purpose: imputation lives inside the model pipeline so that it is fitted on training folds only.",
    )
    miss = ", ".join(f"{k} ({v})" for k, v in eda["missing_by_column"].items())
    table(
        doc,
        ["Property", "Value"],
        [
            ["Rows / features", f"{eda['rows']} / 13 (+ binary target)"],
            [
                "Positive rate",
                f"{eda['positive_rate']:.1%} ({round(eda['positive_rate'] * eda['rows'])} patients with disease)",
            ],
            ["Missing cells", f"{eda['missing_cells']}: {miss} (encoded as '?')"],
            [
                "Strongest monotone correlates (|Spearman rho|)",
                ", ".join(f"{k} {v:.2f}" for k, v in eda["top_correlates_with_target"].items()),
            ],
        ],
        [6, 10.5],
    )
    figure(
        doc,
        FIG / "eda" / "01_class_balance.png",
        "Figure 2. Class balance: near-balanced outcome.",
        7.5,
    )
    figure(
        doc,
        FIG / "eda" / "03_continuous_by_outcome.png",
        "Figure 3. Continuous features by outcome. Patients with disease are older, reach a lower maximum heart rate and show larger ST depression.",
        16.5,
    )
    figure(
        doc,
        FIG / "eda" / "04_prevalence_by_category.png",
        "Figure 4. Disease prevalence by category (dashed = overall rate). Asymptomatic chest pain, reversible thallium defect, exercise angina, flat/down ST slope and more coloured vessels raise the rate sharply.",
        16,
    )
    figure(
        doc,
        FIG / "eda" / "05_correlation_heatmap.png",
        "Figure 5. Spearman correlation. Feature-feature correlations stay below 0.6, so multicollinearity is not a concern.",
        13.5,
    )
    figure(
        doc,
        FIG / "eda" / "06_age_vs_max_heart_rate.png",
        "Figure 6. Chronotropic response: diseased patients fall below the healthy cloud at the same age, motivating the engineered hr_reserve_pct feature.",
        11.5,
    )
    para(
        doc,
        "The executed notebook `notebooks/01_exploratory_analysis.ipynb` reproduces every figure using the same "
        "functions as the `heartsense.explore` module; it is generated from `scripts/build_notebook.py` so it "
        "stays in sync with the code.",
    )

    # ---------------------------------------------------------------- 4. features + models
    doc.add_heading("4. Feature engineering and model development", 1)
    doc.add_heading("4.1 Preprocessing pipeline", 2)
    para(
        doc,
        "A single scikit-learn Pipeline performs every transformation, so the exact fitted steps are replayed "
        "at prediction time and the API contains no preprocessing logic of its own:",
    )
    bullets(
        doc,
        [
            (
                "ClinicalFeatureEngineer: ",
                "adds hr_reserve_pct = max_heart_rate / (220 - age), st_burden = st_depression x "
                "slope weight (up 1.0, flat 1.5, down 2.0), and risk_factor_count (age >= 55, cholesterol > 240, "
                "resting BP >= 140, high fasting sugar, exercise angina).",
            ),
            (
                "Continuous + count + engineered columns: ",
                "median imputation then standardisation.",
            ),
            ("Binary columns: ", "most-frequent imputation, passed through."),
            (
                "Nominal columns (chest_pain_type, rest_ecg, st_slope, thalassemia): ",
                "most-frequent imputation then one-hot "
                "encoding with handle_unknown='ignore' (25 model inputs in total).",
            ),
        ],
    )
    doc.add_heading("4.2 Model families, tuning and selection", 2)
    para(
        doc,
        "Five families were tuned with RandomizedSearchCV (24 sampled configurations each, 12 for k-NN) "
        "under RepeatedStratifiedKFold (5 folds x 2 repeats, seed 2025) on an 80 % stratified training split "
        "(242 patients). Scoring tracked ROC-AUC, accuracy, precision, recall and F1; refit used ROC-AUC. The "
        "champion is the family with the highest mean CV ROC-AUC (ties broken on CV recall, since a missed "
        "diagnosis is the costlier error). The 20 % hold-out split (61 patients) is reported but never used "
        "for selection, so it remains an unbiased estimate.",
    )
    rows = []
    for _, r in board.iterrows():
        rows.append(
            [
                r["model"].replace("_", " "),
                f"{r['cv_roc_auc_mean']:.3f} ± {r['cv_roc_auc_std']:.3f}",
                f"{r['cv_train_roc_auc_mean']:.3f}",
                f"{r['cv_accuracy_mean']:.3f}",
                f"{r['cv_recall_mean']:.3f}",
                f"{r['test_roc_auc']:.3f}",
                f"{r['test_accuracy']:.3f}",
                f"{r['test_recall']:.3f}",
                f"{r['test_precision']:.3f}",
                int(r["test_false_negatives"]),
            ]
        )
    table(
        doc,
        [
            "Model",
            "CV ROC-AUC",
            "CV train AUC",
            "CV acc",
            "CV recall",
            "Test AUC",
            "Test acc",
            "Test recall",
            "Test prec",
            "FN",
        ],
        rows,
        [3.2, 2.4, 1.6, 1.3, 1.5, 1.4, 1.4, 1.6, 1.5, 0.8],
        font_size=8,
    )
    hp = ", ".join(
        f"{k}={v if not isinstance(v, float) else round(v, 4)}"
        for k, v in meta["hyperparameters"].items()
    )
    para(
        doc,
        f"Champion hyper-parameters: {hp}. Logistic regression and the RBF-SVM are statistically "
        "indistinguishable on CV (difference far below one standard deviation); the linear model is preferred "
        "because it has the smallest train/validation gap, trains in milliseconds and its coefficients are "
        "interpretable. The random forest and gradient boosting fit the training folds almost perfectly "
        "(train AUC 0.97-0.98) while validating lower: classic over-fitting on 242 rows.",
    )
    figure(
        doc,
        FIG / "model" / "leaderboard.png",
        "Figure 7. Leaderboard: CV ROC-AUC (with std) next to hold-out ROC-AUC.",
        13,
    )
    figure(
        doc,
        FIG / "model" / "roc_curves.png",
        "Figure 8. Hold-out ROC curves of the five tuned families.",
        10.5,
    )
    figure(
        doc,
        FIG / "model" / "calibration.png",
        "Figure 9. Reliability diagram on the hold-out set.",
        10.5,
    )
    figure(
        doc,
        FIG / "model" / f"{meta['family']}_confusion.png",
        "Figure 10. Champion confusion matrix (threshold 0.5).",
        8,
    )
    figure(
        doc,
        FIG / "model" / f"{meta['family']}_importance.png",
        "Figure 11. Permutation importance of the champion on raw inputs (drop in ROC-AUC when shuffled).",
        11.5,
    )

    # ---------------------------------------------------------------- 5. tracking
    doc.add_heading("5. Experiment tracking with MLflow", 1)
    para(
        doc,
        "MLflow runs with a SQLite backend (`mlflow.db`) and a local artifact store (`mlruns/`), both uploaded "
        "as CI artifacts. The run structure mirrors the training procedure:",
    )
    bullets(
        doc,
        [
            (
                "Family runs (5): ",
                "best hyper-parameters, CV mean/std for every metric, hold-out metrics, confusion matrix, "
                "permutation-importance plot and table, the full cv_results table, the list of transformed "
                "feature names and the fitted pipeline logged with a signature and input example.",
            ),
            (
                "Trial runs (108, nested): ",
                "one per sampled configuration with its parameters, CV scores and rank, so the "
                "search can be inspected with parallel-coordinates or compared across families.",
            ),
            (
                "model-selection run: ",
                "leaderboard table, ROC / PR / calibration overlays, the selection rule as a parameter.",
            ),
            (
                "Model Registry: ",
                f"the champion is registered as `{meta['model_name']}` version {meta['registry_version']} "
                "with the alias `production`, a `family` tag and a description that records the CV score and git SHA.",
            ),
        ],
    )
    figure(
        doc,
        SHOTS / "mlflow_01_experiment_runs.png",
        "Figure 12. MLflow experiment: family runs with nested trials.",
        16,
    )
    figure(
        doc,
        SHOTS / "mlflow_02_champion_run_overview.png",
        "Figure 13. Champion run: parameters, metrics and tags.",
        16,
    )
    figure(
        doc,
        SHOTS / "mlflow_06_selection_leaderboard.png",
        "Figure 14. Artifacts of the model-selection run.",
        16,
    )
    figure(
        doc,
        SHOTS / "mlflow_07_registered_model.png",
        "Figure 15. Model Registry entry with the production alias.",
        16,
    )

    # ---------------------------------------------------------------- 6. packaging
    doc.add_heading("6. Model packaging and reproducibility", 1)
    bullets(
        doc,
        [
            (
                "Artefacts: ",
                "`model_store/heartsense_pipeline.joblib` (the full pipeline) and `model_store/metadata.json` "
                "(model version, family, hyper-parameters, every metric, dataset URL and hashes, split, seed, "
                "decision threshold, risk tiers, input feature list, library versions, git SHA, MLflow run id). "
                "The same model is logged to MLflow with a signature and registered.",
            ),
            (
                "Environments: ",
                "`requirements-api.txt` pins the eight runtime packages used in the image; `requirements.txt` "
                "adds MLflow, plotting, notebook and test tooling, also pinned.",
            ),
            (
                "Determinism: ",
                "SEED = 2025 for the split, the CV folds and every estimator; the raw dataset is pinned by hash; "
                "the clean CSV is committed; the metadata records the clean file's hash too.",
            ),
            (
                "Offline inference: ",
                "`python -m heartsense.predict --input examples/batch_request.json` reuses the exact "
                "scoring function the API calls.",
            ),
        ],
    )

    # ---------------------------------------------------------------- 7. CI/CD + tests
    doc.add_heading("7. CI/CD pipeline and automated testing", 1)
    para(
        doc,
        "The test suite (`tests/`, 56 tests) covers: data acquisition and cleaning (checksum verification, "
        "tampered-file rejection, schema, target derivation, missing-value handling, invalid codes, implausible "
        "readings, purity, de-duplication); feature engineering (formulas, NaN propagation, no mutation, "
        "transformed-matrix shape and finiteness, unseen categories at inference); model selection helpers "
        "(search space, metric computation, threshold behaviour, parameter serialisation); offline prediction; "
        "and the HTTP contract (health, readiness, metadata, high- and low-risk predictions, optional fields, "
        "nine out-of-range cases, unknown fields, batch limits, Prometheus counters, JSON logs without patient "
        "data, 503 when the model is missing). A slow end-to-end test runs the real training loop against a "
        "temporary MLflow store and checks the registry alias.",
    )
    figure(
        doc,
        SHOTS / "tests_01_pytest.png",
        "Figure 16. Local run of the test suite with coverage.",
        15.5,
    )
    figure(doc, SHOTS / "tests_02_ruff.png", "Figure 17. Lint and formatting gate.", 15.5)
    para(
        doc,
        "GitHub Actions workflow `mlops-pipeline.yml` (triggered on push, pull request and manually):",
    )
    table(
        doc,
        ["Job", "What it does", "Artifacts"],
        [
            ["quality", "ruff check (GitHub annotations) and ruff format --check", "lint logs"],
            [
                "helm-lint",
                "helm lint, helm template, kubeconform on the rendered chart and monitoring manifests",
                "rendered manifests",
            ],
            [
                "unit-tests (3.11, 3.12)",
                "ingest dataset, pytest with coverage and JUnit output, coverage summary in the job page",
                "junit, coverage xml, logs",
            ],
            [
                "train-model",
                "EDA + full training; leaderboard and champion written to the job summary",
                "model_store, reports, mlflow.db, mlruns, logs",
            ],
            [
                "container",
                "docker buildx build, run, smoke test; on main push :sha7 and :latest to ghcr.io",
                "build, smoke-test and container logs",
            ],
        ],
        [3.2, 9, 4.3],
    )
    para(
        doc,
        "Jobs are chained with `needs`, so a failing lint, test, training or smoke test stops the pipeline "
        "before anything is published; `concurrency` cancels superseded runs. Step summaries make the "
        "leaderboard and coverage visible without downloading artifacts.",
    )
    figure(
        doc, SHOTS / "github_01_actions_run.png", "Figure 18. Workflow run with the job graph.", 16
    )
    figure(
        doc,
        SHOTS / "github_02_train_summary.png",
        "Figure 19. Train job summary with the leaderboard.",
        16,
    )
    figure(
        doc,
        SHOTS / "github_03_container_job.png",
        "Figure 20. Container job: build, smoke test, publish.",
        16,
    )

    # ---------------------------------------------------------------- 8. docker
    doc.add_heading("8. Model containerisation", 1)
    para(
        doc,
        "The Dockerfile is multi-stage: a builder installs the pinned runtime requirements and the heartsense "
        "package into a virtualenv; the runtime stage copies only that virtualenv, the service code and the "
        "two model files onto python:3.12-slim, creates a system user (uid 10001) and runs uvicorn with one "
        "worker (Prometheus counters are per process, so scaling is done with replicas). A HEALTHCHECK polls "
        "/readyz. `.dockerignore` whitelists exactly what the image needs.",
    )
    para(
        doc,
        "`scripts/smoke_test.sh` exercises a running container end to end: wait for /readyz, /healthz, "
        "/metadata, a high-risk prediction (must be 'disease'), a low-risk prediction (must be 'no_disease'), "
        "a 3-patient batch, an invalid payload (must be 422) and the /metrics endpoint. CI runs the same script.",
    )
    figure(doc, SHOTS / "docker_01_build.png", "Figure 21. Image build.", 15.5)
    figure(
        doc, SHOTS / "docker_02_smoke_test.png", "Figure 22. Container run and smoke test.", 15.5
    )
    table(
        doc,
        ["Endpoint", "Purpose"],
        [
            [
                "POST /predict",
                "one patient -> prediction, risk_score, confidence, risk_tier (low/moderate/elevated/high), model_version, request_id, latency_ms",
            ],
            ["POST /predict/batch", "up to 200 patients"],
            ["GET /healthz, /readyz", "liveness; readiness returns 503 until the model is loaded"],
            ["GET /metadata", "model card (metadata.json)"],
            [
                "GET /metrics",
                "Prometheus exposition: requests, latency histogram, predictions by label and tier, risk-score histogram, validation errors, model_ready, in-flight",
            ],
            ["GET /docs", "Swagger UI"],
        ],
        [4.2, 12.3],
    )

    # ---------------------------------------------------------------- 9. deployment
    doc.add_heading("9. Production deployment on Kubernetes", 1)
    para(
        doc,
        "The API is deployed with a Helm chart (`deploy/helm/heartsense-api`) to a minikube cluster "
        "(Kubernetes v1.37, docker driver). The chart renders a ServiceAccount, a ConfigMap for environment "
        "variables (its checksum is a pod annotation, so config changes roll the pods), a Deployment with two "
        "replicas and a rolling update that never drops below full capacity, startup/readiness/liveness probes, "
        "resource requests and limits, a non-root read-only container security context, a HorizontalPodAutoscaler "
        "(2-5 replicas at 70 % CPU), a PodDisruptionBudget and a Service of type LoadBalancer (an Ingress is "
        "available behind a flag). Pods carry prometheus.io annotations for scraping.",
    )
    code(
        doc,
        "bash scripts/minikube_up.sh\nIMAGE=heartsense-api:1.0.0 bash scripts/deploy.sh      # helm upgrade --install + monitoring\n"
        "bash scripts/port_forward.sh                           # or: minikube -p heartsense tunnel\n"
        "bash scripts/smoke_test.sh http://127.0.0.1:8000",
    )
    figure(
        doc,
        SHOTS / "k8s_01_workloads.png",
        "Figure 23. Helm release, deployment, pods, services, HPA and PDB.",
        15.5,
    )
    figure(
        doc,
        SHOTS / "k8s_02_rollout_and_describe.png",
        "Figure 24. Rollout status and pod details (probes, limits).",
        15.5,
    )
    figure(
        doc,
        SHOTS / "k8s_03_smoke_via_service.png",
        "Figure 25. Smoke test through the Kubernetes Service.",
        15.5,
    )
    figure(
        doc,
        SHOTS / "api_02_predict_try_it_out.png",
        "Figure 26. Swagger UI: live prediction through the cluster.",
        15.5,
    )

    # ---------------------------------------------------------------- 10. monitoring
    doc.add_heading("10. Monitoring and logging", 1)
    para(
        doc,
        "Logging: a middleware emits exactly one JSON line per request (timestamp, level, request_id, method, "
        "path, status, latency_ms) and the prediction handler adds an event with the label and risk tier. "
        "Patient measurements are never logged; a test enforces this. Request ids are accepted from the "
        "X-Request-ID header or generated, and returned in the response for correlation.",
    )
    para(
        doc,
        "Metrics: the service exposes heartsense_http_requests_total, heartsense_http_request_duration_seconds, "
        "heartsense_predictions_total{label,risk_tier}, heartsense_risk_score (histogram), "
        "heartsense_validation_errors_total, heartsense_model_ready, heartsense_model_info and "
        "heartsense_inflight_requests. Prometheus discovers pods in the heartsense namespace through "
        "annotations, so new replicas are scraped automatically. Five alert rules cover: model not ready, "
        "5xx ratio > 5 %, p95 /predict latency > 250 ms, positive-prediction share far above the 46 % training "
        "prevalence (a simple drift signal) and absence of traffic.",
    )
    para(
        doc,
        "Grafana is provisioned from files (datasource, dashboard provider, dashboard JSON) with anonymous "
        "read-only access. The 14-panel dashboard shows replicas scraped, model readiness, request rate, p95, "
        "error ratio, total predictions, rate by path/status, latency quantiles, predictions by risk tier, "
        "positive share vs. training prevalence, a risk-score heatmap, validation rejections, in-flight and "
        "per-pod request rates.",
    )
    figure(
        doc,
        SHOTS / "grafana_01_dashboard_kiosk.png",
        "Figure 27. Grafana dashboard under synthetic traffic.",
        16.5,
    )
    figure(
        doc,
        SHOTS / "prometheus_01_targets.png",
        "Figure 28. Prometheus targets: both replicas UP.",
        15.5,
    )
    figure(doc, SHOTS / "prometheus_02_alert_rules.png", "Figure 29. Alert rules.", 15.5)
    figure(
        doc,
        SHOTS / "k8s_04_json_logs.png",
        "Figure 30. Structured JSON request logs from the pods.",
        15.5,
    )

    # ---------------------------------------------------------------- 11. conclusion
    doc.add_heading("11. Discussion, limitations and next steps", 1)
    bullets(
        doc,
        [
            (
                "Data size: ",
                "303 rows limit what any model can learn; the CV standard deviation (about 0.03 AUC) is the "
                "honest uncertainty. Repeated CV and selecting on CV only keep the reported hold-out numbers credible.",
            ),
            (
                "Threshold: ",
                "the default 0.5 threshold gives recall 0.86; a screening deployment would lower it and accept more "
                "false positives. The threshold is a metadata field, so it can change without retraining.",
            ),
            (
                "Drift: ",
                "the dashboard and alert compare the live positive share with the training prevalence; a next step is "
                "logging feature distributions (without identifiers) and a scheduled retraining workflow.",
            ),
            (
                "Delivery: ",
                "the Helm chart is cloud-ready (LoadBalancer / Ingress, HPA); pointing kubectl at GKE/EKS/AKS and running "
                "`scripts/deploy.sh` with the GHCR image is the only change needed for a public deployment.",
            ),
        ],
    )
    doc.add_heading("12. Repository and links", 1)
    bullets(
        doc,
        [
            ("Code: ", a.repo),
            ("CI runs: ", a.repo + "/actions"),
            ("Container image: ", "ghcr.io/mahamaya-mishra/heartsense-api"),
            ("Video: ", a.video),
            ("Dataset: ", "https://archive.ics.uci.edu/dataset/45/heart+disease (CC BY 4.0)"),
        ],
    )
    para(
        doc,
        "All screenshots referenced in this report are in the repository under screenshots/; the executed EDA "
        "notebook, leaderboard, figures, Helm chart, manifests and workflow are in the folders named above.",
        italic=True,
        size=9.5,
    )

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
